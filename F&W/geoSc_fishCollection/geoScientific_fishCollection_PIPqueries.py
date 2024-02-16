import warnings
warnings.simplefilter(action='ignore')

import os
import json
import pyodbc
import pandas as pd
from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils.dataframe import dataframe_to_rows


def get_db_cnxinfo (dbname='BCGW'):
    """ Retrieves db connection params from the config file"""
    
    with open(r'H:\config\db_config.json', 'r') as file:
        data = json.load(file)
        
    if dbname in data:
        cnxinfo = data[dbname]

        return cnxinfo
    
    raise KeyError(f"Database '{dbname}' not found.")
    
    
def connect_to_DB (driver,server,port,dbq, username,password):
    """ Returns a connection to Oracle database"""
    try:
        connectString =f"""
                    DRIVER={driver};
                    SERVER={server}:{port};
                    DBQ={dbq};
                    Uid={username};
                    Pwd={password}
                       """
        connection = pyodbc.connect(connectString)
        #cursor= connection.cursor()
        print  ("...Successffuly connected to the database")
    except:
        raise Exception('...Connection failed! Please check your connection parameters')

    return connection


def make_xlsx(df_dict, xlsx_path):
    """Exports dataframes to an .xlsx file"""
    # Create a new workbook
    workbook = Workbook()

    # Remove the default "Sheet" created by Workbook
    default_sheet = workbook.get_sheet_by_name('Sheet')
    workbook.remove(default_sheet)

    # Export each DF in dict as sheet within a single XLSX
    for key, df in df_dict.items():
        # Create a worksheet for each DataFrame
        sheet = workbook.create_sheet(title=key)

        # Write the DataFrame to the sheet
        for row in dataframe_to_rows(df, index=False, header=True):
            sheet.append(row)

        # Set the column width dynamically based on the length of the text
        for column in sheet.columns:
            max_length = 0
            column = [cell for cell in column]
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = max(15, min(max_length + 2, 30))
            sheet.column_dimensions[column[0].column_letter].width = adjusted_width

        # Remove spaces from the sheet name for the table name
        table_name = key.replace(' ', '_')

        # Create a table using the data in the sheet
        tab = Table(displayName=table_name, ref=sheet.dimensions)

        # Add a TableStyleInfo to the table
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        tab.tableStyleInfo = style

        # Add the table to the sheet
        sheet.add_table(tab)

    # Save the workbook
    workbook.save(xlsx_path)
    
    
def run_analysis ():
    print ('Connect to BCGW.')
    driver= [x for x in pyodbc.drivers() if x.startswith('Oracle')][0]  
    cnxinfo= get_db_cnxinfo(dbname='BCGW')
    server = cnxinfo['server']
    port= cnxinfo['port']
    dbq= cnxinfo['dbq']
    username= cnxinfo['username']
    password= cnxinfo['password']

    connection= connect_to_DB (driver,server,port,dbq,username,password)

    print ("Execute SQL Query")

    ############################# INPUTS ######################################
    long= 620984
    lat= 6201344
    ############################# INPUTS ######################################

    sql_pip = f"""
    SELECT
        pip.CNSLTN_AREA_NAME AS CONSULTATION_AREA
    FROM
        WHSE_ADMIN_BOUNDARIES.PIP_CONSULTATION_AREAS_SP pip
    WHERE
        SDO_RELATE(pip.SHAPE, SDO_GEOMETRY('POINT({long} {lat})', 26910), 'mask=ANYINTERACT') = 'TRUE'
        """


    sql_soi = f"""
    SELECT
        soi.NAME
    FROM
        REG_LEGAL_AND_ADMIN_BOUNDARIES.QSOI_BC_REGIONS soi
    WHERE
        SDO_RELATE(soi.GEOMETRY, SDO_GEOMETRY('POINT({long} {lat})', 26910), 'mask=ANYINTERACT') = 'TRUE' 
    """
       
    df_pip = pd.read_sql(sql_pip, connection)
    df_soi =pd.read_sql(sql_soi, connection)

    df_pip.drop_duplicates(subset=['CONSULTATION_AREA'], inplace= True)

    df_dict={}  
    df_dict['PIP overlap']= df_pip
    df_dict['SOI overlap']= df_soi

    connection.close()

    print ("Export Report")
    outloc= r'y:/WORKSPACE_2024/20240201_FW_pip_watersheds'
    filename= f'geoSc_fishCollection_FN_report_locUTM10_x{long}_y{lat}.xlsx'
    make_xlsx(df_dict, os.path.join(outloc,filename))
    
run_analysis ()