import os
import pandas as pd
from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils.dataframe import dataframe_to_rows

def make_xlsx(df_dict, xlsx_path):
    """Exports dataframes to an .xlsx file"""
    # Create a new workbook
    workbook = Workbook()

    # Remove the default "Sheet" created by Workbook
    default_sheet = workbook.get_sheet_by_name('Sheet')
    workbook.remove(default_sheet)

    # Export each DF in dict as sheet within a single XLSX
    for key, df in df_dict.items():
        # Create a worksheet for each DataFrame
        sheet = workbook.create_sheet(title=key)

        # Write the DataFrame to the sheet
        for row in dataframe_to_rows(df, index=False, header=True):
            sheet.append(row)

        # Set the column width dynamically based on the length of the text
        for column in sheet.columns:
            max_length = 0
            column = [cell for cell in column]
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = max(15, min(max_length + 2, 30))
            sheet.column_dimensions[column[0].column_letter].width = adjusted_width

        # Remove spaces from the sheet name for the table name
        table_name = key.replace(' ', '_')

        # Create a table using the data in the sheet
        tab = Table(displayName=table_name, ref=sheet.dimensions)

        # Add a TableStyleInfo to the table
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        tab.tableStyleInfo = style

        # Add the table to the sheet
        sheet.add_table(tab)

    # Save the workbook to the specified path
    workbook.save(xlsx_path)
