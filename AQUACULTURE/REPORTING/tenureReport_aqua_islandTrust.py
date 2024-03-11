"""
This script runs SQL queries and export results
in an xlsx report. 

Provide SQLs in the load_queries() function
"""

import warnings
warnings.simplefilter(action='ignore')

import json
import pyodbc
import pandas as pd
from datetime import datetime

from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils.dataframe import dataframe_to_rows

class OracleConnector:
    def __init__(self, dbname='BCGW'):
        self.dbname = dbname
        self.cnxinfo = self.get_db_cnxinfo()

    def get_db_cnxinfo(self):
        """ Retrieves db connection params from the config file"""
        with open(r'H:\config\db_config.json', 'r') as file:
            data = json.load(file)
        
        if self.dbname in data:
            return data[self.dbname]
        
        raise KeyError(f"Database '{self.dbname}' not found.")

    def connect_to_db(self):
        """ Connects to Oracle DB and create a cursor"""
        try:
            driver= [x for x in pyodbc.drivers() if x.startswith('Oracle')][0]  
            self.connection_string =f"""
                        DRIVER={driver};
                        SERVER={self.cnxinfo['server']}:{self.cnxinfo['port']};
                        DBQ={self.cnxinfo['dbq']};
                        Uid={self.cnxinfo['username']};
                        Pwd={self.cnxinfo['password']}
                        """
            self.connection = pyodbc.connect(self.connection_string)
            self.cursor= self.connection.cursor()
            print  ("..Successffuly connected to the database")
        except Exception as e:
            raise Exception('..Connection failed:', e)
            
    def disconnect_db(self):
        """Close the Oracle connection and cursor"""
        if hasattr(self, 'cursor') and self.cursor:
            self.cursor.close()
        if hasattr(self, 'connection') and self.connection:
            self.connection.close()
            print("....Disconnected from the database")


def load_queries ():
    """Returns a dictionnaries of sql queries"""
    sql={}
    sql['Active Tenures']="""
            SELECT* FROM(
            SELECT
                  AD.ADMIN_AREA_NAME,  
                  CAST(IP.INTRID_SID AS NUMBER) INTEREST_PARCEL_ID,
                  CAST(DT.DISPOSITION_TRANSACTION_SID AS NUMBER) DISPOSITION_TRANSACTION_ID,
                  DS.FILE_CHR AS FILE_NBR,
                  SG.STAGE_NME AS STAGE,
                  --TT.ACTIVATION_CDE,
                  TT.STATUS_NME AS STATUS,
                  DT.APPLICATION_TYPE_CDE AS APPLICATION_TYPE,
                  TS.EFFECTIVE_DAT AS EFFECTIVE_DATE,
                  TY.TYPE_NME AS TENURE_TYPE,
                  ST.SUBTYPE_NME AS TENURE_SUBTYPE,
                  PU.PURPOSE_NME AS TENURE_PURPOSE,
                  SP.SUBPURPOSE_NME AS TENURE_SUBPURPOSE,
                  --DT.DOCUMENT_CHR,
                  DT.RECEIVED_DAT AS RECEIVED_DATE,
                  --DT.ENTERED_DAT AS ENTERED_DATE,
                  DT.COMMENCEMENT_DAT AS COMMENCEMENT_DATE,
                  DT.EXPIRY_DAT AS EXPIRY_DATE,
                  --IP.AREA_CALC_CDE,
                  ROUND(IP.AREA_HA_NUM,2) AS AREA_HA,
                  DT.LOCATION_DSC,
                  --OU.UNIT_NAME,
                  --IP.LEGAL_DSC,
                  CONCAT(PR.LEGAL_NAME, PR.FIRST_NAME || ' ' || PR.LAST_NAME) AS CLIENT_NAME_PRIMARY
                  --SH.SHAPE
                  
            FROM WHSE_TANTALIS.TA_DISPOSITION_TRANSACTIONS DT 
              JOIN WHSE_TANTALIS.TA_INTEREST_PARCELS IP 
                ON DT.DISPOSITION_TRANSACTION_SID = IP.DISPOSITION_TRANSACTION_SID
                  AND IP.EXPIRY_DAT IS NULL
              JOIN WHSE_TANTALIS.TA_DISP_TRANS_STATUSES TS
                ON DT.DISPOSITION_TRANSACTION_SID = TS.DISPOSITION_TRANSACTION_SID 
                  AND TS.EXPIRY_DAT IS NULL
              JOIN WHSE_TANTALIS.TA_DISPOSITIONS DS
                ON DS.DISPOSITION_SID = DT.DISPOSITION_SID
              JOIN WHSE_TANTALIS.TA_STAGES SG 
                ON SG.CODE_CHR = TS.CODE_CHR_STAGE
              JOIN WHSE_TANTALIS.TA_STATUS TT 
                ON TT.CODE_CHR = TS.CODE_CHR_STATUS
              JOIN WHSE_TANTALIS.TA_AVAILABLE_TYPES TY 
                ON TY.TYPE_SID = DT.TYPE_SID    
              JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBTYPES ST 
                ON ST.SUBTYPE_SID = DT.SUBTYPE_SID 
                  AND ST.TYPE_SID = DT.TYPE_SID 
              JOIN WHSE_TANTALIS.TA_AVAILABLE_PURPOSES PU 
                ON PU.PURPOSE_SID = DT.PURPOSE_SID    
              JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBPURPOSES SP 
                ON SP.SUBPURPOSE_SID = DT.SUBPURPOSE_SID 
                  AND SP.PURPOSE_SID = DT.PURPOSE_SID 
              JOIN WHSE_TANTALIS.TA_ORGANIZATION_UNITS OU 
                ON OU.ORG_UNIT_SID = DT.ORG_UNIT_SID 
              JOIN WHSE_TANTALIS.TA_TENANTS TE 
                ON TE.DISPOSITION_TRANSACTION_SID = DT.DISPOSITION_TRANSACTION_SID
                  AND TE.SEPARATION_DAT IS NULL
                  AND TE.PRIMARY_CONTACT_YRN = 'Y'
              JOIN WHSE_TANTALIS.TA_INTERESTED_PARTIES PR
                ON PR.INTERESTED_PARTY_SID = TE.INTERESTED_PARTY_SID
            	
              JOIN WHSE_TANTALIS.TA_INTEREST_PARCEL_SHAPES SH
                ON SH.INTRID_SID = IP.INTRID_SID
                
              JOIN WHSE_LEGAL_ADMIN_BOUNDARIES.ABMS_LGL_ADMIN_AREAS_SVW AD 
                ON SDO_RELATE(SH.SHAPE, AD.SHAPE, 'mask=ANYINTERACT')= 'TRUE'
                    AND AD.ADMIN_AREA_TYPE = 'Local Trust Area'
                    AND AD.ADMIN_AREA_GROUP_NAME= 'Islands Trust'
              
              ) TN
            
            WHERE 
                TN.STATUS = 'DISPOSITION IN GOOD STANDING' 
                AND TENURE_PURPOSE =  'AQUACULTURE'
             
            ORDER BY TN.EFFECTIVE_DATE DESC
            """
            
    sql['Expired Tenures']="""
            SELECT* FROM(
            SELECT
                  AD.ADMIN_AREA_NAME,  
                  CAST(IP.INTRID_SID AS NUMBER) INTEREST_PARCEL_ID,
                  CAST(DT.DISPOSITION_TRANSACTION_SID AS NUMBER) DISPOSITION_TRANSACTION_ID,
                  DS.FILE_CHR AS FILE_NBR,
                  SG.STAGE_NME AS STAGE,
                  --TT.ACTIVATION_CDE,
                  TT.STATUS_NME AS STATUS,
                  DT.APPLICATION_TYPE_CDE AS APPLICATION_TYPE,
                  TS.EFFECTIVE_DAT AS EFFECTIVE_DATE,
                  TY.TYPE_NME AS TENURE_TYPE,
                  ST.SUBTYPE_NME AS TENURE_SUBTYPE,
                  PU.PURPOSE_NME AS TENURE_PURPOSE,
                  SP.SUBPURPOSE_NME AS TENURE_SUBPURPOSE,
                  --DT.DOCUMENT_CHR,
                  DT.RECEIVED_DAT AS RECEIVED_DATE,
                  --DT.ENTERED_DAT AS ENTERED_DATE,
                  DT.COMMENCEMENT_DAT AS COMMENCEMENT_DATE,
                  DT.EXPIRY_DAT AS EXPIRY_DATE,
                  --IP.AREA_CALC_CDE,
                  ROUND(IP.AREA_HA_NUM,2) AS AREA_HA,
                  DT.LOCATION_DSC,
                  --OU.UNIT_NAME,
                  --IP.LEGAL_DSC,
                  CONCAT(PR.LEGAL_NAME, PR.FIRST_NAME || ' ' || PR.LAST_NAME) AS CLIENT_NAME_PRIMARY
                  --SH.SHAPE
                  
            FROM WHSE_TANTALIS.TA_DISPOSITION_TRANSACTIONS DT 
              JOIN WHSE_TANTALIS.TA_INTEREST_PARCELS IP 
                ON DT.DISPOSITION_TRANSACTION_SID = IP.DISPOSITION_TRANSACTION_SID
                  AND IP.EXPIRY_DAT IS NULL
              JOIN WHSE_TANTALIS.TA_DISP_TRANS_STATUSES TS
                ON DT.DISPOSITION_TRANSACTION_SID = TS.DISPOSITION_TRANSACTION_SID 
                  AND TS.EXPIRY_DAT IS NULL
              JOIN WHSE_TANTALIS.TA_DISPOSITIONS DS
                ON DS.DISPOSITION_SID = DT.DISPOSITION_SID
              JOIN WHSE_TANTALIS.TA_STAGES SG 
                ON SG.CODE_CHR = TS.CODE_CHR_STAGE
              JOIN WHSE_TANTALIS.TA_STATUS TT 
                ON TT.CODE_CHR = TS.CODE_CHR_STATUS
              JOIN WHSE_TANTALIS.TA_AVAILABLE_TYPES TY 
                ON TY.TYPE_SID = DT.TYPE_SID    
              JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBTYPES ST 
                ON ST.SUBTYPE_SID = DT.SUBTYPE_SID 
                  AND ST.TYPE_SID = DT.TYPE_SID 
              JOIN WHSE_TANTALIS.TA_AVAILABLE_PURPOSES PU 
                ON PU.PURPOSE_SID = DT.PURPOSE_SID    
              JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBPURPOSES SP 
                ON SP.SUBPURPOSE_SID = DT.SUBPURPOSE_SID 
                  AND SP.PURPOSE_SID = DT.PURPOSE_SID 
              JOIN WHSE_TANTALIS.TA_ORGANIZATION_UNITS OU 
                ON OU.ORG_UNIT_SID = DT.ORG_UNIT_SID 
              JOIN WHSE_TANTALIS.TA_TENANTS TE 
                ON TE.DISPOSITION_TRANSACTION_SID = DT.DISPOSITION_TRANSACTION_SID
                  AND TE.SEPARATION_DAT IS NULL
                  AND TE.PRIMARY_CONTACT_YRN = 'Y'
              JOIN WHSE_TANTALIS.TA_INTERESTED_PARTIES PR
                ON PR.INTERESTED_PARTY_SID = TE.INTERESTED_PARTY_SID
                
              JOIN WHSE_TANTALIS.TA_INTEREST_PARCEL_SHAPES SH
                ON SH.INTRID_SID = IP.INTRID_SID
                
              JOIN WHSE_LEGAL_ADMIN_BOUNDARIES.ABMS_LGL_ADMIN_AREAS_SVW AD 
                ON SDO_RELATE(SH.SHAPE, AD.SHAPE, 'mask=ANYINTERACT')= 'TRUE'
                    AND AD.ADMIN_AREA_TYPE = 'Local Trust Area'
                    AND AD.ADMIN_AREA_GROUP_NAME= 'Islands Trust'
              
              ) TN
            
              
            WHERE TN.TENURE_PURPOSE =  'AQUACULTURE'
              AND TN.STATUS = 'EXPIRED'
              
              AND TN.FILE_NBR IN (SELECT CROWN_LANDS_FILE 
                                  FROM WHSE_TANTALIS.TA_CROWN_TENURES_SVW
                                  WHERE RESPONSIBLE_BUSINESS_UNIT = 'VI - LAND MGMNT - VANCOUVER ISLAND SERVICE REGION' 
                                    AND TENURE_STATUS = 'ACCEPTED'
                                    AND APPLICATION_TYPE_CDE = 'REP'
                                    AND CROWN_LANDS_FILE NOT IN (SELECT CROWN_LANDS_FILE 
                                         FROM WHSE_TANTALIS.TA_CROWN_TENURES_SVW
                                         WHERE TENURE_STATUS = 'DISPOSITION IN GOOD STANDING')) 
             
                                            
              AND TN.EXPIRY_DATE = (SELECT MAX(TG.EXPIRY_DATE)
                                    FROM (SELECT
                                              CAST(IP.INTRID_SID AS NUMBER) INTRID_SID,
                                              CAST(DT.DISPOSITION_TRANSACTION_SID AS NUMBER) DISPOSITION_TRANSACTION_SID,
                                              DS.FILE_CHR AS FILE_NBR,
                                              DT.EXPIRY_DAT AS EXPIRY_DATE,
                                              TT.STATUS_NME AS STATUS
                
                                          FROM WHSE_TANTALIS.TA_DISPOSITION_TRANSACTIONS DT 
                                            JOIN WHSE_TANTALIS.TA_INTEREST_PARCELS IP 
                                              ON DT.DISPOSITION_TRANSACTION_SID = IP.DISPOSITION_TRANSACTION_SID
                                                AND IP.EXPIRY_DAT IS NULL
                                            JOIN WHSE_TANTALIS.TA_DISP_TRANS_STATUSES TS
                                              ON DT.DISPOSITION_TRANSACTION_SID = TS.DISPOSITION_TRANSACTION_SID 
                                                AND TS.EXPIRY_DAT IS NULL
                                            JOIN WHSE_TANTALIS.TA_DISPOSITIONS DS
                                              ON DS.DISPOSITION_SID = DT.DISPOSITION_SID
                                            JOIN WHSE_TANTALIS.TA_STAGES SG 
                                              ON SG.CODE_CHR = TS.CODE_CHR_STAGE
                                            JOIN WHSE_TANTALIS.TA_STATUS TT 
                                              ON TT.CODE_CHR = TS.CODE_CHR_STATUS
                                            JOIN WHSE_TANTALIS.TA_AVAILABLE_TYPES TY 
                                              ON TY.TYPE_SID = DT.TYPE_SID    
                                            JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBTYPES ST 
                                              ON ST.SUBTYPE_SID = DT.SUBTYPE_SID 
                                                AND ST.TYPE_SID = DT.TYPE_SID 
                                            JOIN WHSE_TANTALIS.TA_AVAILABLE_PURPOSES PU 
                                              ON PU.PURPOSE_SID = DT.PURPOSE_SID    
                                            JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBPURPOSES SP 
                                              ON SP.SUBPURPOSE_SID = DT.SUBPURPOSE_SID 
                                                AND SP.PURPOSE_SID = DT.PURPOSE_SID 
                                             JOIN WHSE_TANTALIS.TA_ORGANIZATION_UNITS OU 
                                              ON OU.ORG_UNIT_SID = DT.ORG_UNIT_SID)TG
                                              
                                            WHERE TG.STATUS = 'EXPIRED'
                                           AND  TG.FILE_NBR = TN.FILE_NBR
                                            
                                      GROUP BY TG.FILE_NBR)
            
            ORDER BY TN.EXPIRY_DATE DESC
            """        
   
    sql['New Applications']="""
        SELECT* FROM(
        SELECT
              AD.ADMIN_AREA_NAME,  
              CAST(IP.INTRID_SID AS NUMBER) INTEREST_PARCEL_ID,
              CAST(DT.DISPOSITION_TRANSACTION_SID AS NUMBER) DISPOSITION_TRANSACTION_ID,
              DS.FILE_CHR AS FILE_NBR,
              SG.STAGE_NME AS STAGE,
              --TT.ACTIVATION_CDE,
              TT.STATUS_NME AS STATUS,
              DT.APPLICATION_TYPE_CDE AS APPLICATION_TYPE,
              TS.EFFECTIVE_DAT AS EFFECTIVE_DATE,
              TY.TYPE_NME AS TENURE_TYPE,
              ST.SUBTYPE_NME AS TENURE_SUBTYPE,
              PU.PURPOSE_NME AS TENURE_PURPOSE,
              SP.SUBPURPOSE_NME AS TENURE_SUBPURPOSE,
              --DT.DOCUMENT_CHR,
              DT.RECEIVED_DAT AS RECEIVED_DATE,
              --DT.ENTERED_DAT AS ENTERED_DATE,
              DT.COMMENCEMENT_DAT AS COMMENCEMENT_DATE,
              DT.EXPIRY_DAT AS EXPIRY_DATE,
              --IP.AREA_CALC_CDE,
              ROUND(IP.AREA_HA_NUM,2) AS AREA_HA,
              DT.LOCATION_DSC,
              --OU.UNIT_NAME,
              --IP.LEGAL_DSC,
              CONCAT(PR.LEGAL_NAME, PR.FIRST_NAME || ' ' || PR.LAST_NAME) AS CLIENT_NAME_PRIMARY
              --SH.SHAPE
              
        FROM WHSE_TANTALIS.TA_DISPOSITION_TRANSACTIONS DT 
          JOIN WHSE_TANTALIS.TA_INTEREST_PARCELS IP 
            ON DT.DISPOSITION_TRANSACTION_SID = IP.DISPOSITION_TRANSACTION_SID
              AND IP.EXPIRY_DAT IS NULL
          JOIN WHSE_TANTALIS.TA_DISP_TRANS_STATUSES TS
            ON DT.DISPOSITION_TRANSACTION_SID = TS.DISPOSITION_TRANSACTION_SID 
              AND TS.EXPIRY_DAT IS NULL
          JOIN WHSE_TANTALIS.TA_DISPOSITIONS DS
            ON DS.DISPOSITION_SID = DT.DISPOSITION_SID
          JOIN WHSE_TANTALIS.TA_STAGES SG 
            ON SG.CODE_CHR = TS.CODE_CHR_STAGE
          JOIN WHSE_TANTALIS.TA_STATUS TT 
            ON TT.CODE_CHR = TS.CODE_CHR_STATUS
          JOIN WHSE_TANTALIS.TA_AVAILABLE_TYPES TY 
            ON TY.TYPE_SID = DT.TYPE_SID    
          JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBTYPES ST 
            ON ST.SUBTYPE_SID = DT.SUBTYPE_SID 
              AND ST.TYPE_SID = DT.TYPE_SID 
          JOIN WHSE_TANTALIS.TA_AVAILABLE_PURPOSES PU 
            ON PU.PURPOSE_SID = DT.PURPOSE_SID    
          JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBPURPOSES SP 
            ON SP.SUBPURPOSE_SID = DT.SUBPURPOSE_SID 
              AND SP.PURPOSE_SID = DT.PURPOSE_SID 
          JOIN WHSE_TANTALIS.TA_ORGANIZATION_UNITS OU 
            ON OU.ORG_UNIT_SID = DT.ORG_UNIT_SID 
          JOIN WHSE_TANTALIS.TA_TENANTS TE 
            ON TE.DISPOSITION_TRANSACTION_SID = DT.DISPOSITION_TRANSACTION_SID
              AND TE.SEPARATION_DAT IS NULL
              AND TE.PRIMARY_CONTACT_YRN = 'Y'
          JOIN WHSE_TANTALIS.TA_INTERESTED_PARTIES PR
            ON PR.INTERESTED_PARTY_SID = TE.INTERESTED_PARTY_SID
            
          JOIN WHSE_TANTALIS.TA_INTEREST_PARCEL_SHAPES SH
            ON SH.INTRID_SID = IP.INTRID_SID
            
          JOIN WHSE_LEGAL_ADMIN_BOUNDARIES.ABMS_LGL_ADMIN_AREAS_SVW AD 
            ON SDO_RELATE(SH.SHAPE, AD.SHAPE, 'mask=ANYINTERACT')= 'TRUE'
                AND AD.ADMIN_AREA_TYPE = 'Local Trust Area'
                AND AD.ADMIN_AREA_GROUP_NAME= 'Islands Trust'
          
          ) TN
        
        WHERE TN.STAGE = 'APPLICATION'
          AND TN.STATUS = 'ACCEPTED'
          AND TN.TENURE_PURPOSE =  'AQUACULTURE'
          AND TN.APPLICATION_TYPE = 'NEW' 
         
        ORDER BY TN.EFFECTIVE_DATE DESC
            """
            
    return sql            
            
def make_xlsx(df_dict, xlsx_path):
    """Exports dataframes to an .xlsx file"""
    # Create a new workbook
    workbook = Workbook()

    # Remove the default "Sheet" created by Workbook
    default_sheet = workbook.get_sheet_by_name('Sheet')
    workbook.remove(default_sheet)

    # Export each DF in dict as sheet within a single XLSX
    for key, df in df_dict.items():
        # Create a worksheet for each DataFrame
        sheet = workbook.create_sheet(title=key)

        # Write the DataFrame to the sheet
        for row in dataframe_to_rows(df, index=False, header=True):
            sheet.append(row)

        # Set the column width dynamically based on the length of the text
        for column in sheet.columns:
            max_length = 0
            column = [cell for cell in column]
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = max(15, min(max_length + 2, 30))
            sheet.column_dimensions[column[0].column_letter].width = adjusted_width

        # Remove spaces from the sheet name for the table name
        table_name = key.replace(' ', '_')

        # Create a table using the data in the sheet
        tab = Table(displayName=table_name, ref=sheet.dimensions)

        # Add a TableStyleInfo to the table
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        tab.tableStyleInfo = style

        # Add the table to the sheet
        sheet.add_table(tab)

    # Save the workbook to the specified path
    workbook.save(xlsx_path)
    

if __name__ == "__main__":
    print('\nConnect to BCGW')
    oracle_connector = OracleConnector()
    oracle_connector.connect_to_db()
    
    try:
        print('\nRun queries')
        sql =load_queries()
        df_dict={}
        c= 1
        for k, v in sql.items():
            print (f'..query {c} of {len(sql)}: {k}')
            df = pd.read_sql(v, oracle_connector.connection)
            df_dict[k]= df
            
            for col in df.columns:
                if 'DATE' in col:
                    df[col] =  pd.to_datetime(df[col], 
                                              infer_datetime_format=True, 
                                              errors = 'coerce').dt.date
            c+=1

    except Exception as e:
        raise Exception(f"Error occurred: {e}")
    
    finally:
        oracle_connector.disconnect_db()
    
    print ('\nExport the report')
    today= datetime.today().strftime('%Y%m%d')
    outfile= today + '_tenureReport_aqua_islandTrust.xlsx'
    make_xlsx(df_dict, outfile)